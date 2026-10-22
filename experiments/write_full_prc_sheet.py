"""Write results/full_prc_sign/all_experiments.xlsx with one embedded output image per row."""
from __future__ import annotations

import csv
from io import BytesIO
from pathlib import Path

from openpyxl import Workbook
from openpyxl.drawing.image import Image as XLImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from PIL import Image as PILImage

ROOT = Path(__file__).resolve().parents[1]
SRC = ROOT / "results" / "full_prc_sign" / "all_experiments.csv"
IMAGES = ROOT / "results" / "full_prc_sign" / "images"
OUT = ROOT / "results" / "full_prc_sign" / "all_experiments.xlsx"

COLUMNS = [
    "output_image",
    "image_id",
    "guide",
    "target",
    "watermark",
    "invert_prompt",
    "generate_prompt",
    "steps",
    "strength",
    "mix_weight",
    "guidance",
    "t_final",
    "n_steps",
    "channel",
    "sign_agreement_x",
    "sign_agreement_y",
    "detect",
    "decode_ok",
    "exact_message",
    "message_ber",
    "margin",
    "wrong_key_detect",
    "wrong_key_margin",
    "pixel_mse_vs_original",
    "psnr_vs_original",
    "lpips_vs_original",
    "seconds",
]

EXPLANATIONS = {
    "output_image": "The saved generation. Filename is {image_id}__{guide}__{target}__{watermark}.png. oracle_latent and png rows share this picture; the channel only changes how the codeword is read back.",
    "image_id": "Id of the original EDICT photo. imagenet_dog_1, imagenet_cat, church, fountain, or imagenet_lake, from EDICT/experiment_images/.",
    "guide": "Prompt used to invert the real photo before the sign flip. null is the empty string. caption is that photo's description in grid.json.",
    "target": "Prompt used to generate the output. null, caption, small, or style, with the text in grid.json.",
    "watermark": "prc means both inverted latents were replaced by z_i = c_i * |z_i|. nowm leaves the inverted signs unchanged.",
    "invert_prompt": "The exact inversion string. Blank means the empty string.",
    "generate_prompt": "The exact generation string. Blank means the empty string.",
    "steps": "Number of DDIM steps on the Stable Diffusion 1.4 schedule. Here, 50.",
    "strength": "Edit strength s. 1 uses every one of the 50 steps.",
    "mix_weight": "EDICT coupling weight p. Here, 0.93.",
    "guidance": "Classifier-free guidance scale. With an empty prompt the conditional and unconditional embeddings are the same, so guidance does not change the noise prediction.",
    "t_final": "Noisiest timestep visited. The 50-step grid is every 20th training index, 0, 20, ..., 980. Training indices run 0 through 999, so 980 is the top of this grid.",
    "n_steps": "How many timesteps were in the schedule. Here, 50.",
    "channel": "Where the scored latent came from. oracle_latent inverts the latent pair before the VAE writes pixels. png reloads this file, VAE-encodes it, and inverts. image_only means no recovery was run.",
    "sign_agreement_x": "Fraction of the 16384 coordinates in the recovered x latent whose sign equals the embedded codeword. sign(0) counts as a mismatch.",
    "sign_agreement_y": "Same fraction on the recovered y latent. The codeword was written into both streams.",
    "detect": "Gunn PRC detector on soft signs from x, variance 1.5, FPR 1e-5. True when margin >= 0.",
    "decode_ok": "Gunn decoder returned a payload (belief propagation, then the linear solve, and the test bits matched).",
    "exact_message": "The first 64 bits of that payload equal the sent 64-bit message. False when decode returned nothing.",
    "message_ber": "Fraction of those 64 bits that differ from the sent message. Blank when decode returned nothing.",
    "margin": "True-key detector statistic minus its threshold. Positive means detect.",
    "wrong_key_detect": "The same soft signs scored with a different PRC key.",
    "wrong_key_margin": "Margin for that wrong key.",
    "pixel_mse_vs_original": "Mean squared error versus the source photo after mapping both RGB images to [-1, 1]. 0 is identical. Same definition as EDICT Table 1, where 50-step reconstruction is about 0.015.",
    "psnr_vs_original": "10*log10(4/MSE), peak value 2 because pixels are in [-1, 1]. Higher is closer to the photo. EDICT's 0.015 reconstruction MSE is about 24 dB on this scale. EDICT did not publish PSNR.",
    "lpips_vs_original": "AlexNet LPIPS versus the source photo, both images resized to 256. Lower is closer. EDICT's mammal-edit average against the original photo is 0.27. This is not the PRC paper's pairwise diversity LPIPS.",
    "seconds": "Wall time from the start of this generation until the row was written. A png row includes the oracle inversion written just before it.",
}

NUMBER_COLUMNS = {
    "steps",
    "strength",
    "mix_weight",
    "guidance",
    "t_final",
    "n_steps",
    "sign_agreement_x",
    "sign_agreement_y",
    "message_ber",
    "margin",
    "wrong_key_margin",
    "pixel_mse_vs_original",
    "psnr_vs_original",
    "lpips_vs_original",
    "seconds",
}

THUMB = 112


def thumb(path: Path) -> BytesIO:
    image = PILImage.open(path).convert("RGB")
    image.thumbnail((THUMB, THUMB))
    buf = BytesIO()
    image.save(buf, format="PNG")
    buf.seek(0)
    return buf


def as_cell(name: str, value: str):
    if value == "":
        return None
    if name in NUMBER_COLUMNS:
        return float(value) if "." in value else int(value)
    return value


def main() -> None:
    with SRC.open(newline="") as f:
        rows = list(csv.DictReader(f))

    wb = Workbook()
    ws = wb.active
    ws.title = "experiments"
    header_fill = PatternFill("solid", fgColor="D9E2F3")
    explain_fill = PatternFill("solid", fgColor="FFF2CC")
    wrap = Alignment(wrap_text=True, vertical="top")
    header_font = Font(bold=True)

    for col, name in enumerate(COLUMNS, start=1):
        exp = ws.cell(1, col, EXPLANATIONS[name])
        exp.alignment = wrap
        exp.fill = explain_fill
        exp.font = Font(size=9)
        head = ws.cell(2, col, name)
        head.font = header_font
        head.fill = header_fill
        head.alignment = Alignment(vertical="center", wrap_text=True)

    missing = []
    buffers = []
    for i, row in enumerate(rows):
        excel_row = i + 3
        ws.row_dimensions[excel_row].height = 88
        name = f"{row['image_id']}__{row['guide']}__{row['target']}__{row['watermark']}.png"
        path = IMAGES / name
        if path.exists():
            buf = thumb(path)
            buffers.append(buf)
            img = XLImage(buf)
            img.width = THUMB
            img.height = THUMB
            img.anchor = f"A{excel_row}"
            ws.add_image(img)
        else:
            missing.append(name)
            ws.cell(excel_row, 1, f"missing {name}")
        for col, key in enumerate(COLUMNS[1:], start=2):
            cell = ws.cell(excel_row, col, as_cell(key, row.get(key, "")))
            if key in ("invert_prompt", "generate_prompt"):
                cell.alignment = Alignment(wrap_text=True, vertical="center")

    ws.row_dimensions[1].height = 78
    ws.row_dimensions[2].height = 22
    widths = {
        "output_image": 18,
        "image_id": 18,
        "guide": 12,
        "target": 12,
        "watermark": 12,
        "invert_prompt": 42,
        "generate_prompt": 42,
        "channel": 16,
    }
    for col, name in enumerate(COLUMNS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = widths.get(name, 16)
    ws.freeze_panes = "B3"
    ws.auto_filter.ref = f"A2:{get_column_letter(len(COLUMNS))}{len(rows) + 2}"
    wb.save(OUT)
    print(f"wrote {OUT} rows={len(rows)} missing={len(missing)}")
    for name in missing:
        print(" missing", name)


if __name__ == "__main__":
    main()
